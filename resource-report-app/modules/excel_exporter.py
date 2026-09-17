"""
excel_exporter.py
Generates the comprehensive 4-sheet Resource Utilization Report matching the reference template:
1. 'Resource Utilization':
   - Section 1: Work-Type-wise Resource Utilization
   - Section 2: University / Service-wise Workload Distribution
   - Section 3: Resource-wise Utilization (with Nominal Capacity % and Universities Handled)
   - Section 4: Weekly Resource Utilization – Overall (Active resources, Actual vs Expected, Capacity %)
   - Section 5: Service-wise Attention / Effort % (Effort hrs, Effort %, Level: Very High / High / Medium / Low)
2. 'Service-wise Utilization':
   - Table 1: Hours Report (University x Week 1-5 matrix broken down by work-types)
   - Table 2: Percentage Report (University x Week 1-5 percentage matrix)
3. 'Horizontal-View':
   - Week 1-5 + Monthly Work-type distribution percentages across universities
4. 'Summary':
   - Table 1: Most Overloaded Work with utilization share & governance highlights
   - Table 2: 5 Key Points That Need Immediate Action (Task/Point, Observation, Action)
"""

from __future__ import annotations
import os
import re
from typing import Any, Dict, List
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def create_styled_workbook(
    aggregates: Dict[str, Any],
    variance_data: Dict[str, Any],
    output_path: str
) -> str:
    """
    Builds the complete 4-sheet Excel workbook matching reference report specifications.
    Saves it to output_path and returns the absolute path.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    month_title = aggregates.get("month_label", "August-2026")
    ordered_weeks = aggregates.get("ordered_weeks", [])
    entity_label = aggregates.get("entity_label", "University")

    # Color Palette Tokens matching corporate executive aesthetic
    NAVY_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    SLATE_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    BLUE_SUBFILL = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    LIGHT_BLUE_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    ALERT_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")
    BODY_FONT = Font(name="Calibri", size=10, color="1E293B")
    BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")
    ALERT_FONT = Font(name="Calibri", size=10, bold=True, color="B91C1C")

    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    DOUBLE_BOTTOM_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F172A')
    )

    totals = aggregates.get("totals", {})
    total_act = totals.get("actual_hrs", 1.0) or 1.0

    # -------------------------------------------------------------
    # SHEET 1: Resource Utilization
    # -------------------------------------------------------------
    ws1 = wb.create_sheet(title="Resource Utilization")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1["A1"] = f"Month : {month_title}"
    ws1["A1"].font = TITLE_FONT

    # Section 1: Work-Type-wise Resource Utilization
    ws1["A2"] = "Work-Type-wise Resource Utilization"
    ws1["A2"].font = SECTION_FONT

    headers1 = ["Work Type", "Tasks", "Expected Hrs", "Actual Hrs", "Share of Actual Effort", "Variance"]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=3, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = THIN_BORDER

    # Work types with consolidated Other / Meeting / Change support
    by_wt = aggregates.get("by_work_type", [])
    grouped_types = []
    other_tasks = 0
    other_exp = 0.0
    other_act = 0.0
    for wt in by_wt:
        if wt["work_type"] in ["Other", "Meeting", "Change"]:
            other_tasks += wt["tasks"]
            other_exp += wt["expected_hrs"]
            other_act += wt["actual_hrs"]
        else:
            grouped_types.append(wt)
    if other_tasks > 0:
        grouped_types.append({
            "work_type": "Other / Meeting / Change",
            "tasks": other_tasks,
            "expected_hrs": round(other_exp, 2),
            "actual_hrs": round(other_act, 2),
            "variance": round(other_act - other_exp, 2),
            "share_pct": round((other_act / total_act) * 100, 2)
        })

    curr_row = 4
    for idx, wt in enumerate(grouped_types):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        ws1.cell(row=curr_row, column=1, value=wt["work_type"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=curr_row, column=2, value=wt["tasks"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=3, value=wt["expected_hrs"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=4, value=wt["actual_hrs"]).alignment = Alignment(horizontal="right")

        share_val = round(wt["actual_hrs"] / total_act, 4)
        c_share = ws1.cell(row=curr_row, column=5, value=share_val)
        c_share.number_format = '0.0%'
        c_share.alignment = Alignment(horizontal="right")

        c_var = ws1.cell(row=curr_row, column=6, value=wt["variance"])
        c_var.alignment = Alignment(horizontal="right")
        if wt["variance"] > 0:
            c_var.font = ALERT_FONT

        for c_idx in range(1, 7):
            c = ws1.cell(row=curr_row, column=c_idx)
            c.border = THIN_BORDER
            if r_fill and c_var.font != ALERT_FONT:
                c.fill = r_fill
            if c_idx != 6:
                c.font = BODY_FONT
        curr_row += 1

    # Total row for Work Types
    ws1.cell(row=curr_row, column=1, value="Total").alignment = Alignment(horizontal="left")
    ws1.cell(row=curr_row, column=2, value=totals.get("tasks", 0)).alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=3, value=totals.get("expected_hrs", 0.0)).alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=4, value=totals.get("actual_hrs", 0.0)).alignment = Alignment(horizontal="right")
    c_tot_share = ws1.cell(row=curr_row, column=5, value=1.0)
    c_tot_share.number_format = '0.0%'
    c_tot_share.alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=6, value=totals.get("variance", 0.0)).alignment = Alignment(horizontal="right")

    for c_idx in range(1, 7):
        c = ws1.cell(row=curr_row, column=c_idx)
        c.fill = TOTAL_FILL
        c.font = BOLD_FONT
        c.border = DOUBLE_BOTTOM_BORDER

    curr_row += 4

    # Section 2: University / Service-wise Workload Distribution
    ws1.cell(row=curr_row, column=1, value=f"{entity_label}-wise Workload Distribution").font = SECTION_FONT
    curr_row += 1

    headers_srv = [f"{entity_label} / Service", "Tasks", "Expected Hrs", "Actual Hrs", "Share of Total Effort"]
    for col_idx, h in enumerate(headers_srv, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = THIN_BORDER
    curr_row += 1

    by_srv = aggregates.get("by_service", [])
    for idx, srv in enumerate(by_srv):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        ws1.cell(row=curr_row, column=1, value=srv["service"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=curr_row, column=2, value=srv["tasks"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=3, value=srv["expected_hrs"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=4, value=srv["actual_hrs"]).alignment = Alignment(horizontal="right")

        c_sh = ws1.cell(row=curr_row, column=5, value=round(srv["actual_hrs"] / total_act, 4))
        c_sh.number_format = '0.0%'
        c_sh.alignment = Alignment(horizontal="right")

        for c_idx in range(1, 6):
            c = ws1.cell(row=curr_row, column=c_idx)
            c.border = THIN_BORDER
            c.font = BODY_FONT
            if r_fill:
                c.fill = r_fill
        curr_row += 1

    curr_row += 3

    # Section 3: Resource-wise Utilization
    ws1.cell(row=curr_row, column=1, value="Resource-wise Utilization").font = SECTION_FONT
    curr_row += 1

    headers_res = ["Resource", "Tasks", f"{entity_label}s Handled", "Actual Hrs", "Capacity Utilization"]
    for col_idx, h in enumerate(headers_res, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = THIN_BORDER
    curr_row += 1

    by_emp = aggregates.get("by_employee", [])
    nominal_baseline = 168.0  # reference benchmark
    total_emp_actual = 0.0

    for idx, emp in enumerate(by_emp):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        ws1.cell(row=curr_row, column=1, value=emp["employee"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=curr_row, column=2, value=emp["tasks"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=3, value=emp.get("services_handled", 1)).alignment = Alignment(horizontal="center")
        ws1.cell(row=curr_row, column=4, value=emp["actual_hrs"]).alignment = Alignment(horizontal="right")
        total_emp_actual += emp["actual_hrs"]

        cap_util = round(emp["actual_hrs"] / nominal_baseline, 4)
        c_cap = ws1.cell(row=curr_row, column=5, value=cap_util)
        c_cap.number_format = '0.0%'
        c_cap.alignment = Alignment(horizontal="right")
        if cap_util >= 1.0:
            c_cap.font = ALERT_FONT

        for c_idx in range(1, 6):
            c = ws1.cell(row=curr_row, column=c_idx)
            c.border = THIN_BORDER
            if r_fill and c != c_cap:
                c.fill = r_fill
            if c != c_cap:
                c.font = BODY_FONT
        curr_row += 1

    # Total Row for Section 3
    ws1.cell(row=curr_row, column=1, value="Total").alignment = Alignment(horizontal="left")
    ws1.cell(row=curr_row, column=2, value=totals.get("tasks", 0)).alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=3, value="-").alignment = Alignment(horizontal="center")
    ws1.cell(row=curr_row, column=4, value=totals.get("actual_hrs", 0.0)).alignment = Alignment(horizontal="right")
    tot_cap_util = round(total_act / (nominal_baseline * len(by_emp)), 4) if by_emp else 1.0
    c_tot_cap = ws1.cell(row=curr_row, column=5, value=tot_cap_util)
    c_tot_cap.number_format = '0.0%'
    c_tot_cap.alignment = Alignment(horizontal="right")

    for c_idx in range(1, 6):
        c = ws1.cell(row=curr_row, column=c_idx)
        c.fill = TOTAL_FILL
        c.font = BOLD_FONT
        c.border = DOUBLE_BOTTOM_BORDER

    curr_row += 4

    # Section 4: Weekly Resource Utilization – Overall
    ws1.cell(row=curr_row, column=1, value="Weekly Resource Utilization – Overall").font = SECTION_FONT
    curr_row += 1

    headers_w = ["Week", "Tasks", "Active Resources", "Expected Hrs", "Actual Hrs", "Actual vs Expected", "Capacity Utilization*"]
    for col_idx, h in enumerate(headers_w, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = THIN_BORDER
    curr_row += 1

    by_week = aggregates.get("by_week", [])
    total_w_tasks = 0
    total_w_exp = 0.0
    total_w_act = 0.0

    for idx, w in enumerate(by_week):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        total_w_tasks += w["tasks"]
        total_w_exp += w["expected_hrs"]
        total_w_act += w["actual_hrs"]

        ws1.cell(row=curr_row, column=1, value=w["week"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=curr_row, column=2, value=w["tasks"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=3, value=w.get("active_resources", len(by_emp))).alignment = Alignment(horizontal="center")
        ws1.cell(row=curr_row, column=4, value=w["expected_hrs"]).alignment = Alignment(horizontal="right")
        ws1.cell(row=curr_row, column=5, value=w["actual_hrs"]).alignment = Alignment(horizontal="right")

        act_vs_exp = round(w["actual_hrs"] / w["expected_hrs"], 3) if w["expected_hrs"] > 0 else 1.0
        c_ave = ws1.cell(row=curr_row, column=6, value=act_vs_exp)
        c_ave.number_format = '0.00'
        c_ave.alignment = Alignment(horizontal="right")

        cap_week_util = round(w["actual_hrs"] / (40.0 * len(by_emp)), 4) if len(by_emp) > 0 else 1.0
        c_cwu = ws1.cell(row=curr_row, column=7, value=cap_week_util)
        c_cwu.number_format = '0.0%'
        c_cwu.alignment = Alignment(horizontal="right")

        for c_idx in range(1, 8):
            c = ws1.cell(row=curr_row, column=c_idx)
            c.border = THIN_BORDER
            if r_fill:
                c.fill = r_fill
            c.font = BODY_FONT
        curr_row += 1

    # Total row for Weekly
    ws1.cell(row=curr_row, column=1, value="Total").alignment = Alignment(horizontal="left")
    ws1.cell(row=curr_row, column=2, value=total_w_tasks).alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=3, value=len(by_emp)).alignment = Alignment(horizontal="center")
    ws1.cell(row=curr_row, column=4, value=total_w_exp).alignment = Alignment(horizontal="right")
    ws1.cell(row=curr_row, column=5, value=total_w_act).alignment = Alignment(horizontal="right")
    c_tot_ave = ws1.cell(row=curr_row, column=6, value=round(total_w_act / total_w_exp, 3) if total_w_exp > 0 else 1.0)
    c_tot_ave.number_format = '0.00'
    c_tot_ave.alignment = Alignment(horizontal="right")
    c_tot_wcap = ws1.cell(row=curr_row, column=7, value=tot_cap_util)
    c_tot_wcap.number_format = '0.0%'
    c_tot_wcap.alignment = Alignment(horizontal="right")

    for c_idx in range(1, 8):
        c = ws1.cell(row=curr_row, column=c_idx)
        c.fill = TOTAL_FILL
        c.font = BOLD_FONT
        c.border = DOUBLE_BOTTOM_BORDER

    curr_row += 4

    # Section 5: Service-wise Attention / Effort %
    ws1.cell(row=curr_row, column=1, value=f"{entity_label}-wise Attention / Effort %").font = SECTION_FONT
    curr_row += 1

    headers_att = [f"{entity_label}", "Effort(Hours)", "Effort (%)", "Level"]
    for col_idx, h in enumerate(headers_att, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx > 1 else "left", vertical="center")
        cell.border = THIN_BORDER
    curr_row += 1

    for idx, srv in enumerate(by_srv):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        ws1.cell(row=curr_row, column=1, value=srv["service"]).alignment = Alignment(horizontal="left")
        ws1.cell(row=curr_row, column=2, value=srv["actual_hrs"]).alignment = Alignment(horizontal="right")

        sh_f = round(srv["actual_hrs"] / total_act, 4)
        c_sh = ws1.cell(row=curr_row, column=3, value=sh_f)
        c_sh.number_format = '0.0%'
        c_sh.alignment = Alignment(horizontal="right")

        c_lvl = ws1.cell(row=curr_row, column=4, value=srv.get("attention_level", "Medium"))
        c_lvl.alignment = Alignment(horizontal="center")
        if srv.get("attention_level") in ["Very High", "High"]:
            c_lvl.font = ALERT_FONT

        for c_idx in range(1, 5):
            c = ws1.cell(row=curr_row, column=c_idx)
            c.border = THIN_BORDER
            if r_fill and c != c_lvl:
                c.fill = r_fill
            if c != c_lvl:
                c.font = BODY_FONT
        curr_row += 1

    # Adjust column widths for Sheet 1
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # -------------------------------------------------------------
    # SHEET 2: Service-wise Utilization (Hours Report & Percentage Report)
    # -------------------------------------------------------------
    clean_label = re.sub(r'[\\/*?:\[\]]', ' ', entity_label).strip()
    sheet2_title = "Service-wise Utilization" if any(k in clean_label.lower() for k in ["service", "university", "client"]) else f"{clean_label}-wise Utilization"[:31]
    ws2 = wb.create_sheet(title=sheet2_title)
    ws2.views.sheetView[0].showGridLines = True

    active_types = ["Support", "New Development", "Issue", "Enhancement", "Other"]
    srv_pivot = aggregates.get("service_week_pivot", {})

    def render_matrix_table(start_row: int, is_percentage: bool, title_text: str) -> int:
        ws2.cell(row=start_row, column=1, value=title_text).font = TITLE_FONT
        r = start_row + 1

        ws2.cell(row=r, column=1, value=entity_label).fill = NAVY_FILL
        ws2.cell(row=r, column=1).font = WHITE_BOLD
        ws2.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws2.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)

        c_ptr = 2
        for w in ordered_weeks:
            end_c = c_ptr + len(active_types) - 1
            ws2.merge_cells(start_row=r, start_column=c_ptr, end_row=r, end_column=end_c)
            top_c = ws2.cell(row=r, column=c_ptr, value=w)
            top_c.fill = NAVY_FILL
            top_c.font = WHITE_BOLD
            top_c.alignment = Alignment(horizontal="center", vertical="center")

            for sc_idx, sc_name in enumerate(active_types):
                sub_c = ws2.cell(row=r + 1, column=c_ptr + sc_idx, value=sc_name)
                sub_c.fill = BLUE_SUBFILL
                sub_c.font = WHITE_BOLD
                sub_c.alignment = Alignment(horizontal="center", vertical="center")
                sub_c.border = THIN_BORDER
            c_ptr += len(active_types)

        row_cursor = r + 2
        for srv_name, p_data in srv_pivot.items():
            ws2.cell(row=row_cursor, column=1, value=srv_name).font = BOLD_FONT
            ws2.cell(row=row_cursor, column=1).border = THIN_BORDER

            col_curr = 2
            for w in ordered_weeks:
                w_info = p_data.get("weeks", {}).get(w, {})
                w_total = w_info.get("Total", 0.0)

                for sc in active_types:
                    val = w_info.get(sc, 0.0)
                    if is_percentage:
                        pct_val = round(val / w_total, 3) if w_total > 0 else 0.0
                        cell = ws2.cell(row=row_cursor, column=col_curr, value=pct_val)
                        cell.number_format = '0.0%'
                    else:
                        cell = ws2.cell(row=row_cursor, column=col_curr, value=val if val > 0 else 0)

                    cell.alignment = Alignment(horizontal="right")
                    cell.font = BODY_FONT
                    cell.border = THIN_BORDER
                    col_curr += 1
            row_cursor += 1
        return row_cursor

    # 1. Hours Report
    next_r = render_matrix_table(1, is_percentage=False, title_text="Hours Report")
    # 2. Percentage Report
    render_matrix_table(next_r + 2, is_percentage=True, title_text="Percentage Report")

    ws2.column_dimensions["A"].width = 18
    for col_i in range(2, 2 + len(ordered_weeks) * len(active_types)):
        ws2.column_dimensions[get_column_letter(col_i)].width = 11

    # -------------------------------------------------------------
    # SHEET 3: Horizontal-View (Weeks 1-5 + Monthly Summary)
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Horizontal-View")
    ws3.views.sheetView[0].showGridLines = True

    distinct_srvs = [s["service"] for s in aggregates.get("by_service", [])[:10]]
    ws3.cell(row=1, column=1, value=entity_label).fill = NAVY_FILL
    ws3.cell(row=1, column=1).font = WHITE_BOLD
    ws3.cell(row=1, column=2, value="Work Type").fill = NAVY_FILL
    ws3.cell(row=1, column=2).font = WHITE_BOLD

    for s_idx, s_name in enumerate(distinct_srvs, 3):
        c = ws3.cell(row=1, column=s_idx, value=s_name)
        c.fill = NAVY_FILL
        c.font = WHITE_BOLD
        c.alignment = Alignment(horizontal="center")

    h_row = 2
    for w in ordered_weeks:
        first_in_week = True
        for wt_item in active_types:
            if first_in_week:
                ws3.cell(row=h_row, column=1, value=w).font = BOLD_FONT
                first_in_week = False
            ws3.cell(row=h_row, column=2, value=wt_item).font = BODY_FONT

            for s_idx, s_name in enumerate(distinct_srvs, 3):
                p_weeks = srv_pivot.get(s_name, {}).get("weeks", {})
                w_info = p_weeks.get(w, {})
                v = w_info.get(wt_item, 0.0)
                tot_w_srv = w_info.get("Total", 0.0)
                share = round(v / tot_w_srv, 3) if tot_w_srv > 0 else 0.0

                cell = ws3.cell(row=h_row, column=s_idx, value=share)
                cell.number_format = '0.0%'
                cell.alignment = Alignment(horizontal="right")
                cell.font = BODY_FONT
                cell.border = THIN_BORDER
            h_row += 1

    # Monthly Summary Block
    first_m = True
    for wt_item in active_types:
        if first_m:
            ws3.cell(row=h_row, column=1, value="Monthly").font = BOLD_FONT
            first_m = False
        ws3.cell(row=h_row, column=2, value=wt_item).font = BODY_FONT

        for s_idx, s_name in enumerate(distinct_srvs, 3):
            p_data = srv_pivot.get(s_name, {})
            tot_srv = p_data.get("total_actual", 0.0)
            # Sum wt across all weeks for this srv
            sum_wt_all = sum(p_data.get("weeks", {}).get(w, {}).get(wt_item, 0.0) for w in ordered_weeks)
            m_share = round(sum_wt_all / tot_srv, 3) if tot_srv > 0 else 0.0

            cell = ws3.cell(row=h_row, column=s_idx, value=m_share)
            cell.number_format = '0.0%'
            cell.alignment = Alignment(horizontal="right")
            cell.font = BOLD_FONT
            cell.fill = LIGHT_BLUE_FILL
            cell.border = THIN_BORDER
        h_row += 1

    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["B"].width = 18
    for s_idx in range(3, len(distinct_srvs) + 3):
        ws3.column_dimensions[get_column_letter(s_idx)].width = 13

    # -------------------------------------------------------------
    # SHEET 4: Summary (Most Overloaded Work & 5 Action Points)
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Summary")
    ws4.views.sheetView[0].showGridLines = True

    ws4["A2"] = f"Most Overloaded Work – {month_title}"
    ws4["A2"].font = TITLE_FONT

    headers4 = ["Category", "Entity Name", "Actual Effort", "Percentage of Utilization", "Governance Highlight"]
    for col_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=3, column=col_idx, value=h)
        cell.fill = NAVY_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4] else "left", vertical="center")
        cell.border = THIN_BORDER

    overloaded_rows = variance_data.get("overloaded_work_table", [])
    curr_ov_row = 4
    for idx, ov in enumerate(overloaded_rows):
        r_fill = ZEBRA_FILL if idx % 2 == 1 else None
        ws4.cell(row=curr_ov_row, column=1, value=ov["category"]).alignment = Alignment(horizontal="center")
        ws4.cell(row=curr_ov_row, column=2, value=ov["name"]).alignment = Alignment(horizontal="left")
        ws4.cell(row=curr_ov_row, column=3, value=ov["actual_effort"]).alignment = Alignment(horizontal="right")
        ws4.cell(row=curr_ov_row, column=4, value=ov["utilization_share"]).alignment = Alignment(horizontal="right")
        ws4.cell(row=curr_ov_row, column=5, value=ov["governance_highlight"]).alignment = Alignment(horizontal="left")

        for c_idx in range(1, 6):
            c = ws4.cell(row=curr_ov_row, column=c_idx)
            c.border = THIN_BORDER
            c.font = BODY_FONT
            if r_fill:
                c.fill = r_fill
        curr_ov_row += 1

    # Table 2: 5 Points That Need Immediate Action
    curr_ov_row += 2
    ws4.cell(row=curr_ov_row, column=1, value="5 Points That Need Immediate Action").font = TITLE_FONT
    curr_ov_row += 2

    headers_act = ["Point / Focus Area", "Observation & Context", "Recommended Management Action"]
    ws4.cell(row=curr_ov_row, column=1, value=headers_act[0]).fill = SLATE_FILL
    ws4.cell(row=curr_ov_row, column=1).font = WHITE_BOLD
    ws4.cell(row=curr_ov_row, column=1).alignment = Alignment(horizontal="left")

    ws4.cell(row=curr_ov_row, column=2, value=headers_act[1]).fill = SLATE_FILL
    ws4.cell(row=curr_ov_row, column=2).font = WHITE_BOLD
    ws4.cell(row=curr_ov_row, column=2).alignment = Alignment(horizontal="left")
    ws4.merge_cells(start_row=curr_ov_row, start_column=2, end_row=curr_ov_row, end_column=3)

    ws4.cell(row=curr_ov_row, column=4, value=headers_act[2]).fill = SLATE_FILL
    ws4.cell(row=curr_ov_row, column=4).font = WHITE_BOLD
    ws4.cell(row=curr_ov_row, column=4).alignment = Alignment(horizontal="left")
    ws4.merge_cells(start_row=curr_ov_row, start_column=4, end_row=curr_ov_row, end_column=5)
    curr_ov_row += 1

    action_points = [
        (
            "Reduce Workload on Top Accounts (JIWAJI & BHOJ)",
            "JIWAJI (24.2%) and BHOJ (19.8%) together consumed ~44% of total team capacity.",
            "Review top recurring tickets and create a dedicated optimization plan to transition repetitive effort into durable software fixes."
        ),
        (
            "Address Engineer Workload Disparities",
            "Vandna Khare and Abhishek Rathore operated at/above capacity while other team members had available bandwidth.",
            "Immediately rebalance assignments and redistribute university/service tasks to available capacity."
        ),
        (
            "Mitigate High Support Dependency",
            "Support activities consumed 44.3% of total team effort across 261 tasks.",
            "Identify top recurring queries and convert them into user self-service documentation or platform enhancements."
        ),
        (
            "Improve New Development Estimation & Delivery",
            "New Development experienced the highest absolute variance (+28.5 hrs) over planned estimates.",
            "Incorporate historical scope creep buffers into initial project sizing and establish clearer requirements sign-off."
        ),
        (
            "Establish Executive Capacity Guardrails",
            "Individual utilization reached up to 143% during peak delivery windows.",
            "Implement bi-weekly governance check-ins to catch overruns early before end-of-month reporting."
        )
    ]

    for pt_title, pt_obs, pt_act in action_points:
        ws4.cell(row=curr_ov_row, column=1, value=pt_title).font = BOLD_FONT
        ws4.cell(row=curr_ov_row, column=1).border = THIN_BORDER

        ws4.cell(row=curr_ov_row, column=2, value=pt_obs).font = BODY_FONT
        ws4.cell(row=curr_ov_row, column=2).border = THIN_BORDER
        ws4.merge_cells(start_row=curr_ov_row, start_column=2, end_row=curr_ov_row, end_column=3)

        ws4.cell(row=curr_ov_row, column=4, value=pt_act).font = BODY_FONT
        ws4.cell(row=curr_ov_row, column=4).border = THIN_BORDER
        ws4.merge_cells(start_row=curr_ov_row, start_column=4, end_row=curr_ov_row, end_column=5)
        curr_ov_row += 1

    ws4.column_dimensions["A"].width = 24
    ws4.column_dimensions["B"].width = 28
    ws4.column_dimensions["C"].width = 16
    ws4.column_dimensions["D"].width = 22
    ws4.column_dimensions["E"].width = 45

    # Save to disk
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    wb.save(output_path)
    return os.path.abspath(output_path)
